import os
import sys
import tempfile
import shutil
from pathlib import Path

# Add dewdrop root to sys.path
DEWDROP_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(DEWDROP_ROOT))

from tests.run_cleanroom_verification import (
    TOOL_SPECS,
    execute_tool,
    has_ffmpeg,
)
from PIL import Image, ImageDraw
import pypdf
import docx

def create_sample_files(td: Path):
    samples = {}
    
    # 1. Sample Image (PNG)
    png_p = td / "sample.png"
    im = Image.new("RGBA", (200, 200), (40, 140, 240, 255))
    d = ImageDraw.Draw(im)
    d.rectangle([40, 40, 160, 160], fill=(255, 200, 50, 255))
    im.save(png_p, "PNG")
    samples["png"] = png_p

    # 2. Sample Image (JPEG)
    jpg_p = td / "sample.jpg"
    im_rgb = Image.new("RGB", (200, 200), (200, 80, 80))
    im_rgb.save(jpg_p, "JPEG")
    samples["jpg"] = jpg_p

    # 3. Sample PDF
    pdf_p = td / "sample.pdf"
    writer = pypdf.PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    writer.write(pdf_p)
    samples["pdf"] = pdf_p

    # Multi-page PDF
    pdf2_p = td / "multipage.pdf"
    w2 = pypdf.PdfWriter()
    w2.add_blank_page(width=300, height=300)
    w2.add_blank_page(width=300, height=300)
    w2.write(pdf2_p)
    samples["pdf_multi"] = pdf2_p

    # PDF with text
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas as rl_canvas
    pdf_text_p = td / "text_sample.pdf"
    c = rl_canvas.Canvas(str(pdf_text_p), pagesize=A4)
    c.drawString(100, 700, "Dewdrop Cleanroom Verification Heading")
    c.drawString(100, 680, "This is a paragraph of sample text inside PDF.")
    c.drawString(100, 660, "- Bullet item 1")
    c.drawString(100, 640, "- Bullet item 2")
    c.save()
    samples["pdf_text"] = pdf_text_p

    # 4. Sample Word DOCX
    docx_p = td / "sample.docx"
    doc = docx.Document()
    doc.add_heading("Dewdrop Test Document", level=1)
    doc.add_paragraph("Hello world paragraph from Dewdrop verification.")
    doc.save(str(docx_p))
    samples["docx"] = docx_p

    # 5. Sample Text
    txt_p = td / "sample.txt"
    txt_p.write_text("Hello world text\r\nSecond line\r\n", encoding="utf-8")
    samples["txt"] = txt_p

    # 6. Sample JSON
    json_p = td / "sample.json"
    json_p.write_text('[{"id": 1, "name": "Alpha"}, {"id": 2, "name": "Beta"}]', encoding="utf-8")
    samples["json"] = json_p

    # 7. Sample CSV
    csv_p = td / "sample.csv"
    csv_p.write_text("id,name\n1,Alpha\n2,Beta\n", encoding="utf-8")
    samples["csv"] = csv_p

    # 8. Sample SRT
    srt_p = td / "sample.srt"
    srt_p.write_text("1\n00:00:01,000 --> 00:00:03,000\nHello world\n\n", encoding="utf-8")
    samples["srt"] = srt_p

    # 9. Sample VTT
    vtt_p = td / "sample.vtt"
    vtt_p.write_text("WEBVTT\n\n1\n00:00:01.000 --> 00:00:03.000\nHello world\n\n", encoding="utf-8")
    samples["vtt"] = vtt_p

    # 10. Sample Folder with files and subfolders
    folder_p = td / "sample_folder"
    folder_p.mkdir(exist_ok=True)
    (folder_p / "doc1.txt").write_text("Folder file 1 content", encoding="utf-8")
    (folder_p / "doc2.txt").write_text("Folder file 2 content", encoding="utf-8")
    sub_dir = folder_p / "sub_assets"
    sub_dir.mkdir(exist_ok=True)
    (sub_dir / "nested.json").write_text('{"nested": true}', encoding="utf-8")
    samples["folder"] = folder_p

    # 11. Sample Media (MP4 and MP3 via ffmpeg if available)
    if has_ffmpeg():
        mp4_p = td / "sample.mp4"
        import subprocess
        # Generate 1-second synthetic test video
        subprocess.run([
            "ffmpeg", "-y", "-f", "lavfi", "-i", "testsrc=duration=1:size=160x120:rate=10",
            "-f", "lavfi", "-i", "sine=duration=1:frequency=440",
            "-c:v", "libx264", "-c:a", "aac", "-pix_fmt", "yuv420p",
            str(mp4_p)
        ], capture_output=True)
        samples["mp4"] = mp4_p

        mp3_p = td / "sample.mp3"
        subprocess.run([
            "ffmpeg", "-y", "-f", "lavfi", "-i", "sine=duration=1:frequency=440",
            "-c:a", "libmp3lame", str(mp3_p)
        ], capture_output=True)
        samples["mp3"] = mp3_p

    return samples

def run_audit():
    td = Path(tempfile.mkdtemp(prefix="dewdrop_audit_"))
    try:
        samples = create_sample_files(td)
        print(f"Sample files prepared at: {td}")
        
        results = {}
        for tid, spec in TOOL_SPECS.items():
            try:
                # Prepare inputs based on tool
                inp = None
                opts = {}
                
                # Image tools
                if tid in ("img_to_png", "img_to_gif", "img_to_ico"):
                    inp = [samples["jpg"]]
                elif tid in ("img_to_jpeg", "img_to_webp"):
                    inp = [samples["png"]]
                elif tid == "img_compress":
                    inp = [samples["jpg"]]
                    opts = {"quality": 40, "target_kb": 200}
                elif tid == "img_resize":
                    inp = [samples["png"]]
                    opts = {"resize_mode": "percent", "percent": 50}
                elif tid == "img_crop":
                    inp = [samples["png"]]
                    opts = {"ratio": "1:1"}
                elif tid == "img_rotate_flip":
                    inp = [samples["png"]]
                    opts = {"rotation": 90, "flip": "h"}
                elif tid == "img_effects":
                    inp = [samples["png"]]
                    opts = {"effect": "grayscale"}
                elif tid == "img_watermark":
                    inp = [samples["png"]]
                    opts = {"text": "VERIFIED"}
                elif tid == "img_strip_meta":
                    inp = [samples["jpg"]]
                
                # PDF tools
                elif tid == "pdf_from_images":
                    inp = [samples["png"], samples["jpg"]]
                    opts = {"combine": True, "page": "A4"}
                elif tid == "pdf_merge":
                    inp = [samples["pdf"], samples["pdf_multi"]]
                elif tid == "pdf_split":
                    inp = [samples["pdf_multi"]]
                elif tid in ("pdf_extract_pages", "pdf_delete_pages"):
                    inp = [samples["pdf_multi"]]
                    opts = {"pages": "1"}
                elif tid == "pdf_reverse":
                    inp = [samples["pdf_multi"]]
                elif tid == "pdf_rotate":
                    inp = [samples["pdf_multi"]]
                    opts = {"rotation": 90}
                elif tid == "pdf_compress":
                    inp = [samples["pdf_text"]]
                    opts = {"level": 2}
                elif tid == "pdf_to_text":
                    inp = [samples["pdf_text"]]
                elif tid == "pdf_extract_images":
                    inp = [samples["pdf_text"]]
                elif tid == "pdf_strip_meta":
                    inp = [samples["pdf"]]
                elif tid == "pdf_to_word":
                    inp = [samples["pdf_text"]]
                elif tid == "word_to_pdf":
                    inp = [samples["docx"]]

                # Media tools
                elif tid in ("media_to_mp4", "media_to_webm", "video_compress", "video_to_gif", "video_mute", "video_frame_png"):
                    if "mp4" in samples:
                        inp = [samples["mp4"]]
                        opts = {"level": 2, "maxh": 720, "start": "0", "duration": 1}
                elif tid in ("media_to_mp3", "media_to_m4a", "media_to_wav", "media_to_flac", "audio_compress"):
                    if "mp3" in samples:
                        inp = [samples["mp3"]]
                        opts = {"bitrate": 128}
                elif tid == "media_trim":
                    if "mp4" in samples:
                        inp = [samples["mp4"]]
                        opts = {"start": "00:00", "end": "00:01"}

                # Archive tools
                elif tid == "archive_zip":
                    inp = [samples["txt"], samples["json"]]
                    opts = {"level": 6, "combine": True}
                elif tid == "archive_targz":
                    inp = [samples["txt"], samples["json"]]
                    opts = {"level": 6}
                elif tid == "archive_gzip":
                    inp = [samples["txt"]]
                    opts = {"level": 6}
                elif tid == "archive_extract":
                    # Create a test zip first
                    test_zip = td / "audit_extract_test.zip"
                    import zipfile
                    with zipfile.ZipFile(test_zip, "w") as zf:
                        zf.writestr("test.txt", "extract test")
                    inp = [test_zip]

                # Data & text tools
                elif tid in ("json_pretty", "json_minify", "json_to_csv"):
                    inp = [samples["json"]]
                elif tid == "csv_to_json":
                    inp = [samples["csv"]]
                elif tid in ("text_utf8", "text_crlf"):
                    inp = [samples["txt"]]
                elif tid == "srt_to_vtt":
                    inp = [samples["srt"]]
                elif tid == "vtt_to_srt":
                    inp = [samples["vtt"]]
                elif tid == "checksums_all":
                    inp = [samples["txt"]]

                if inp is None:
                    results[tid] = ("SKIP", "Missing sample dependencies (e.g. FFmpeg)")
                    continue

                out_p = execute_tool(tid, inp, opts)
                if out_p and Path(out_p).exists() and (Path(out_p).stat().st_size > 0 or Path(out_p).is_dir()):
                    size = Path(out_p).stat().st_size if Path(out_p).is_file() else sum(f.stat().st_size for f in Path(out_p).glob("**/*") if f.is_file())
                    results[tid] = ("PASS", f"Output {Path(out_p).name} ({size} bytes)")
                else:
                    results[tid] = ("FAIL", f"Output missing or empty: {out_p}")
            except Exception as e:
                results[tid] = ("FAIL", f"Exception: {type(e).__name__}: {e}")

        # Also test Folder Zipping specifically
        try:
            folder_zip_out = execute_tool("archive_zip", [samples["folder"]], {"level": 6})
            import zipfile
            with zipfile.ZipFile(folder_zip_out, "r") as zf:
                names = zf.namelist()
            if len(names) > 1 and any("doc1.txt" in n for n in names):
                results["archive_zip_FOLDER"] = ("PASS", f"Packed {len(names)} items from folder: {names}")
            else:
                results["archive_zip_FOLDER"] = ("FAIL", f"Folder zip only has {len(names)} entries: {names}")
        except Exception as e:
            results["archive_zip_FOLDER"] = ("FAIL", f"Exception: {e}")

        # Also test Folder TAR.GZ specifically
        try:
            folder_tar_out = execute_tool("archive_targz", [samples["folder"]], {"level": 6})
            import tarfile
            with tarfile.open(folder_tar_out, "r:gz") as tf:
                names = tf.getnames()
            if len(names) > 1 and any("doc1.txt" in n for n in names):
                results["archive_targz_FOLDER"] = ("PASS", f"Packed {len(names)} items from folder: {names}")
            else:
                results["archive_targz_FOLDER"] = ("FAIL", f"Folder tar only has {len(names)} entries: {names}")
        except Exception as e:
            results["archive_targz_FOLDER"] = ("FAIL", f"Exception: {e}")

        print("\n" + "="*80)
        print("COMPLETE 50 TOOLS AUDIT RESULTS:")
        print("="*80)
        passed = 0
        failed = 0
        for tid, (st, det) in sorted(results.items()):
            symbol = "✓" if st == "PASS" else ("✗" if st == "FAIL" else "○")
            print(f"[{symbol}] {tid:25} : {st} - {det}")
            if st == "PASS":
                passed += 1
            elif st == "FAIL":
                failed += 1
        print("="*80)
        print(f"SUMMARY: {passed} PASSED, {failed} FAILED, {len(results)} TOTAL TESTED")
        print("="*80)

    finally:
        shutil.rmtree(td, ignore_errors=True)

if __name__ == "__main__":
    run_audit()
